import csv
from pathlib import Path

from openpyxl import load_workbook

from backend.algor.binary_search import binary_search_exact, lower_bound, upper_bound
from backend.algor.merge_sort import merge_sort
from backend.data_struct.hash_table import HashTable
from backend.models import InventoryItem


class InventoryCatalog:
    def __init__(self):
        self._table: HashTable[tuple[str, str], InventoryItem] = HashTable()
        self._sorted: list[InventoryItem] = []
        self._type_ranges: dict[str, tuple[int, int]] = {}

    def __len__(self) -> int:
        return len(self._table)

    def all_items(self) -> list[InventoryItem]:
        return list(self._sorted)

    def types(self) -> list[str]:
        return sorted(self._type_ranges.keys())


    def upsert(self, item: InventoryItem) -> None:
        self._table.insert(item.key, item)
        self._rebuild_sorted()

    def bulk_load(self, items: list[InventoryItem], replace: bool = True) -> int:
        if replace:
            self._table = HashTable()
        for item in items:
            self._table.insert(item.key, item)
        self._rebuild_sorted()
        return len(items)

    def load_csv(self, path: str | Path, replace: bool = True) -> int:
        if replace:
            self._table = HashTable()
        with open(path, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            count = 0
            for row in reader:
                item = InventoryItem(
                    type=row["type"].strip(),
                    item_name=row["item_name"].strip(),
                    live_qty=int(row["live_qty"]),
                    reserved_qty=int(row.get("reserved_qty") or 0),
                )
                self._table.insert(item.key, item)
                count += 1
        self._rebuild_sorted()
        return count

    def load_xlsx(
        self,
        path: str | Path,
        sheet_name: str = "Simple-2026",
        replace: bool = True,
    ) -> int:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"worksheet not found: {sheet_name}")
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if headers is None:
                raise ValueError("inventory workbook is empty")

            header_index = {
                str(value).strip(): index
                for index, value in enumerate(headers)
                if value is not None and str(value).strip()
            }
            required = ["Serial ID", "Item Name (New)", "Type", "Live QTY"]
            missing = [name for name in required if name not in header_index]
            if missing:
                raise ValueError(f"missing inventory columns: {', '.join(missing)}")

            merged: dict[tuple[str, str], int] = {}
            for row in rows:
                serial = row[header_index["Serial ID"]]
                item_name = row[header_index["Item Name (New)"]]
                type_ = row[header_index["Type"]]
                qty = row[header_index["Live QTY"]]

                if not serial or not item_name or not type_:
                    continue

                type_text = str(type_).strip()
                item_text = str(item_name).strip()
                if not type_text or not item_text:
                    continue

                if isinstance(qty, bool):
                    continue
                if isinstance(qty, (int, float)):
                    live_qty = int(qty)
                elif isinstance(qty, str) and qty.strip().lower() == "bulk":
                    live_qty = 100
                else:
                    continue

                if live_qty < 0:
                    continue
                key = (type_text, item_text)
                merged[key] = merged.get(key, 0) + live_qty

            if replace:
                self._table = HashTable()
            for (type_text, item_text), live_qty in merged.items():
                self._table.insert(
                    (type_text, item_text),
                    InventoryItem(type_text, item_text, live_qty, reserved_qty=0),
                )
            self._rebuild_sorted()
            return len(merged)
        finally:
            workbook.close()


    def get(self, type_: str, item_name: str) -> InventoryItem | None:
        return self._table.search((type_, item_name))

    def reserve(self, type_: str, item_name: str, qty: int) -> None:
        item = self._require(type_, item_name)
        if qty <= 0 or qty > item.available_qty:
            raise ValueError(
                f"cannot reserve {qty} of {item_name}: {item.available_qty} available"
            )
        item.reserved_qty += qty

    def release_reservation(self, type_: str, item_name: str, qty: int) -> None:
        item = self._require(type_, item_name)
        item.reserved_qty = max(0, item.reserved_qty - qty)

    def issue(self, type_: str, item_name: str, qty: int) -> None:
        item = self._require(type_, item_name)
        item.live_qty -= qty
        item.reserved_qty = max(0, item.reserved_qty - qty)

    def restock_good(self, type_: str, item_name: str, qty: int) -> None:
        item = self._require(type_, item_name)
        item.live_qty += qty

    def _require(self, type_: str, item_name: str) -> InventoryItem:
        item = self._table.search((type_, item_name))
        if item is None:
            raise KeyError(f"no such inventory item: {type_}/{item_name}")
        return item


    def find_exact(self, type_: str, item_name: str) -> InventoryItem | None:
        target = (type_, item_name.casefold(), item_name)
        idx = binary_search_exact(
            self._sorted, target,
            key_func=lambda i: (i.type, i.item_name.casefold(), i.item_name),
        )
        return None if idx == -1 else self._sorted[idx]

    def search_by_name_prefix(
        self, prefix: str, type_: str | None = None
    ) -> list[InventoryItem]:
        prefix_lower = prefix.strip().casefold()
        if type_ and type_ not in self._type_ranges:
            return []
        ranges = (
            [self._type_ranges[type_]] if type_ else list(self._type_ranges.values())
        )
        results: list[InventoryItem] = []
        for start, end in ranges:
            segment = self._sorted[start:end]
            if prefix_lower == "":
                results.extend(segment)
                continue
            lo = lower_bound(segment, prefix_lower, key_func=lambda i: i.item_name.casefold())
            hi = upper_bound(
                segment, prefix_lower + chr(0x10ffff), key_func=lambda i: i.item_name.casefold()
            )
            results.extend(segment[lo:hi])
        return results


    def _rebuild_sorted(self) -> None:
        self._sorted = merge_sort(
            list(self._table.values()),
            key_func=lambda i: (i.type, i.item_name.casefold(), i.item_name),
        )
        self._type_ranges = {}
        start = 0
        for i in range(1, len(self._sorted) + 1):
            at_end = i == len(self._sorted)
            if at_end or self._sorted[i].type != self._sorted[start].type:
                self._type_ranges[self._sorted[start].type] = (start, i)
                start = i


