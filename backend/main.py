"""FastAPI + Jinja2/HTMX web app (section 8 of the brief).

Single-process demo app: one Allocator/Database live in app.state, built
once at startup (rebuilt from SQLite if data already exists, otherwise
seeded from the real RAI XLSX inventory workbook). Role is a simple cookie set by /login (section 13
default: two hard-coded roles, no real auth).
"""

import os
from contextlib import asynccontextmanager
from datetime import date, datetime
from io import BytesIO
from typing import Optional

from fastapi import Cookie, Depends, FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.algor.merge_sort import merge_sort
from backend.allocator import Allocator
from backend.db.database import Database
from backend.inventory import InventoryCatalog
from backend.models import (
    ItemCondition,
    ItemStatus,
    RequestLineItem,
    ReturnLineItem,
)
from backend.notifications import notifier_from_env

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.environ.get("RAI_DB_PATH", os.path.join(BASE_DIR, "backend", "db", "app.db"))
RAI_XLSX = os.environ.get("RAI_XLSX", os.path.join(BASE_DIR, "backend", "db", "Database_inven_RAI.xlsx"))


@asynccontextmanager
async def lifespan(app: FastAPI):
    db = Database(DB_PATH)
    notifier = notifier_from_env()
    if db.load_inventory_items():
        allocator = Allocator.from_db(db, notifier=notifier)
    else:
        catalog = InventoryCatalog()
        catalog.load_xlsx(RAI_XLSX)
        db.save_inventory_items(catalog.all_items())
        allocator = Allocator(catalog, notifier=notifier, db=db)
    app.state.db = db
    app.state.allocator = allocator
    yield
    db.close()


app = FastAPI(title="RAI Student Project Item Allocator", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=os.path.join(BASE_DIR, "frontend", "static")), name="static")
templates = Jinja2Templates(directory=os.path.join(BASE_DIR, "frontend", "templates"))


def get_allocator(request: Request) -> Allocator:
    return request.app.state.allocator


def get_db(request: Request) -> Database:
    return request.app.state.db


def get_role(role: Optional[str] = Cookie(default=None)) -> str:
    return role if role in ("student", "admin") else "student"


def require_admin(role: str = Depends(get_role)) -> str:
    if role != "admin":
        raise HTTPException(status_code=403, detail="Admin role required")
    return role


def priority_label(pick_up_date: date, today: date | None = None) -> str:
    today = today or date.today()
    days = (pick_up_date - today).days
    if days <= 1:
        return "High"
    if days <= 3:
        return "Medium"
    return "Low"


def ctx(request: Request, role: str, **extra) -> dict:
    return {"role": role, "active_nav": extra.pop("active_nav", ""), **extra}


# -- auth / role switch --------------------------------------------------


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, role: str = Depends(get_role)):
    return templates.TemplateResponse(request, "login.html", ctx(request, role, active_nav="login"))


@app.post("/login")
def do_login(chosen_role: str = Form(...)):
    if chosen_role not in ("student", "admin"):
        raise HTTPException(status_code=400, detail="invalid role")
    target = "/dashboard" if chosen_role == "admin" else "/new-request"
    response = RedirectResponse(url=target, status_code=303)
    response.set_cookie("role", chosen_role, max_age=60 * 60 * 24 * 30)
    return response


@app.get("/")
def root():
    return RedirectResponse(url="/login")


# -- dashboard -------------------------------------------------------------


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(
    request: Request,
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
    db: Database = Depends(get_db),
):
    if role != "admin":
        return RedirectResponse(url="/new-request", status_code=303)
    items = alloc.catalog.all_items()
    available_types = sum(1 for i in items if i.status == ItemStatus.AVAILABLE)
    out_of_stock_items = merge_sort(
        [i for i in items if i.status == ItemStatus.OUT_OF_STOCK],
        key_func=lambda i: i.key,
    )

    today = date.today()
    ordered_requests = alloc.admin_ordered_view()
    annotated_requests = [(req, priority_label(req.pick_up_date, today)) for req in ordered_requests]

    return templates.TemplateResponse(
        request,
        "dashboard.html",
        ctx(
            request,
            role,
            active_nav="dashboard",
            available_types=available_types,
            pending_count=len(alloc.pending_heap),
            reserved_count=len(alloc.reserved),
            borrowed_count=len(alloc.borrowed),
            out_of_stock_items=out_of_stock_items[:15],
            total_items=len(items),
            remaining_stock=sum(i.available_qty for i in items),
            top_borrowed=db.top_borrowed_items(limit=3),
            report_date=today.isoformat(),
            rows=annotated_requests,
        ),
    )


# -- daily activity reports ---------------------------------------------------


def _parse_report_day(date_value: str) -> date:
    if not date_value:
        return date.today()
    try:
        return date.fromisoformat(date_value)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid report date")


def _daily_report_data(db: Database, report_day: date) -> tuple[dict, list, list]:
    raw_events = db.load_activity_for_date(report_day)
    labels = {
        "REQUEST_SUBMITTED": "Request submitted",
        "RESERVATION_READY": "Reservation ready",
        "BORROW_CONFIRMED": "Borrow confirmed",
        "RETURN_PROCESSED": "Return processed",
        "SHORTAGE_NOTICE": "Shortage notice sent",
    }
    events = [
        {
            "time": datetime.fromisoformat(row["event_at"]).strftime("%H:%M:%S"),
            "event_type": row["event_type"],
            "label": labels.get(row["event_type"], row["event_type"].replace("_", " ").title()),
            "project_id": row["project_id"],
            "details": row["details"],
        }
        for row in raw_events
    ]
    counts = {
        "submitted": sum(1 for row in raw_events if row["event_type"] == "REQUEST_SUBMITTED"),
        "reserved": sum(1 for row in raw_events if row["event_type"] == "RESERVATION_READY"),
        "borrowed": sum(1 for row in raw_events if row["event_type"] == "BORROW_CONFIRMED"),
        "returned": sum(1 for row in raw_events if row["event_type"] == "RETURN_PROCESSED"),
        "notices": sum(1 for row in raw_events if row["event_type"] == "SHORTAGE_NOTICE"),
        "damaged_units": db.daily_damage_units(report_day),
    }
    top_borrowed = list(db.top_borrowed_items(limit=3, day=report_day))
    return counts, top_borrowed, events


def _build_daily_summary_docx(
    report_day: date,
    counts: dict,
    top_borrowed: list,
    events: list,
) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)

    title = document.add_heading("Daily Activity Summary", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph = document.add_paragraph(report_day.strftime("%d %B %Y"))
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("Daily Totals", level=1)
    totals = document.add_table(rows=1, cols=2)
    totals.style = "Table Grid"
    totals.rows[0].cells[0].text = "Activity"
    totals.rows[0].cells[1].text = "Count"
    summary_rows = [
        ("Requests Submitted", counts["submitted"]),
        ("Reservations Ready", counts["reserved"]),
        ("Borrow Confirmations", counts["borrowed"]),
        ("Returns Processed", counts["returned"]),
        ("Shortage Notices", counts["notices"]),
        ("Damaged Units", counts["damaged_units"]),
    ]
    for label, value in summary_rows:
        cells = totals.add_row().cells
        cells[0].text = str(label)
        cells[1].text = str(value)

    document.add_paragraph()
    document.add_heading("Top 3 Borrowed Items", level=1)
    if top_borrowed:
        top_table = document.add_table(rows=1, cols=4)
        top_table.style = "Table Grid"
        headers = ["Rank", "Item", "Type", "Units Borrowed"]
        for index, label in enumerate(headers):
            top_table.rows[0].cells[index].text = label
        for rank, item in enumerate(top_borrowed, start=1):
            cells = top_table.add_row().cells
            cells[0].text = f"#{rank}"
            cells[1].text = str(item["item_name"])
            cells[2].text = str(item["type"])
            cells[3].text = str(item["total_borrowed"])
    else:
        document.add_paragraph("No items were borrowed on this date.")

    document.add_paragraph()
    document.add_heading("Activity Timeline", level=1)
    if events:
        timeline = document.add_table(rows=1, cols=4)
        timeline.style = "Table Grid"
        headers = ["Time", "Activity", "Project", "Details"]
        for index, label in enumerate(headers):
            timeline.rows[0].cells[index].text = label
        for event in events:
            cells = timeline.add_row().cells
            cells[0].text = event["time"]
            cells[1].text = event["label"]
            cells[2].text = event["project_id"] or "-"
            cells[3].text = event["details"] or "-"
    else:
        document.add_paragraph("No activity was recorded for this date.")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


@app.get("/admin/daily-summary", response_class=HTMLResponse)
def daily_summary_page(
    request: Request,
    date_value: str = "",
    role: str = Depends(require_admin),
    db: Database = Depends(get_db),
):
    report_day = _parse_report_day(date_value)
    counts, top_borrowed, events = _daily_report_data(db, report_day)
    return templates.TemplateResponse(
        request,
        "daily_summary.html",
        ctx(
            request,
            role,
            active_nav="reports",
            report_date=report_day.isoformat(),
            counts=counts,
            top_borrowed=top_borrowed,
            events=events,
        ),
    )


@app.get("/admin/daily-summary.docx")
def download_daily_summary_word(
    date_value: str = "",
    role: str = Depends(require_admin),
    db: Database = Depends(get_db),
):
    report_day = _parse_report_day(date_value)
    counts, top_borrowed, events = _daily_report_data(db, report_day)
    content = _build_daily_summary_docx(report_day, counts, top_borrowed, events)
    filename = f"RAI_Daily_Activity_Summary_{report_day.isoformat()}.docx"
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# -- new request -------------------------------------------------------------


@app.get("/new-request", response_class=HTMLResponse)
def new_request_page(request: Request, role: str = Depends(get_role), alloc: Allocator = Depends(get_allocator)):
    if role == "admin":
        return RedirectResponse(url="/dashboard", status_code=303)
    return templates.TemplateResponse(
        request,
        "new_request.html",
        ctx(request, role, active_nav="new_request", types=alloc.catalog.types(), today=date.today().isoformat()),
    )


@app.get("/components/search", response_class=HTMLResponse)
def search_components(
    request: Request,
    q: str = "",
    type: str = "",
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
):
    # Selecting a Type should immediately show the components/tools that
    # belong to that type, even when the search box is empty.
    if q.strip() or type:
        results = alloc.catalog.search_by_name_prefix(
            q,
            type_=type or None,
        )
    else:
        results = []
    return templates.TemplateResponse(
        request,
        "_partials/component_results.html",
        {"results": results[:100]},
    )


@app.post("/requests")
def create_request(
    request: Request,
    project_name: str = Form(...),
    student_name: str = Form(...),
    student_id: str = Form(...),
    group: str = Form(...),
    pick_up_date: str = Form(...),
    student_email: str = Form(...),
    item_types: list[str] = Form(default=[]),
    item_names: list[str] = Form(default=[]),
    item_quantities: list[str] = Form(default=[]),
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
):
    if role != "student":
        raise HTTPException(status_code=403, detail="Student role required")
    if not student_id.strip():
        raise HTTPException(status_code=400, detail="Student ID is required")
    if not student_email.strip():
        raise HTTPException(status_code=400, detail="Email is required")
    if "@" not in student_email or "." not in student_email.rsplit("@", 1)[-1]:
        raise HTTPException(status_code=400, detail="Enter a valid email address")
    if not item_types:
        raise HTTPException(status_code=400, detail="At least one component/tool is required")
    try:
        pud = date.fromisoformat(pick_up_date)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid pick-up date")

    lines = []
    for t, n, q in zip(item_types, item_names, item_quantities):
        qty = int(q)
        if qty <= 0:
            raise HTTPException(status_code=400, detail=f"Quantity for {n} must be positive")
        lines.append(RequestLineItem(t, n, qty))

    req = alloc.submit_request(
        project_name=project_name,
        student_name=student_name,
        group=group,
        pick_up_date=pud,
        items=lines,
        student_email=student_email,
        student_id=student_id.strip(),
    )
    return RedirectResponse(url=f"/requests/{req.project_id}", status_code=303)


@app.get("/requests/{project_id}", response_class=HTMLResponse)
def request_status(
    project_id: str, request: Request, role: str = Depends(get_role), alloc: Allocator = Depends(get_allocator)
):
    req = alloc.all_requests.search(project_id)
    if req is None:
        raise HTTPException(status_code=404, detail="Request not found")
    return templates.TemplateResponse(
        request,
        "request_status.html", ctx(request, role, active_nav="new_request", req=req)
    )


# -- inventory ---------------------------------------------------------------


@app.get("/inventory", response_class=HTMLResponse)
def inventory_page(
    request: Request,
    q: str = "",
    type: str = "",
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
):
    items = alloc.catalog.search_by_name_prefix(q, type_=type or None) if q.strip() else alloc.catalog.all_items()
    if type and not q.strip():
        items = [i for i in items if i.type == type]
    return templates.TemplateResponse(
        request,
        "inventory.html",
        ctx(request, role, active_nav="inventory", items=items, types=alloc.catalog.types(), q=q, selected_type=type),
    )


@app.get("/inventory/search", response_class=HTMLResponse)
def inventory_search(
    request: Request,
    q: str = "",
    type: str = "",
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
):
    # Same behavior as the New Request selector: choosing a Type immediately
    # shows every inventory item that belongs to that Type. Typing a name
    # narrows the results inside the selected Type.
    if q.strip():
        items = alloc.catalog.search_by_name_prefix(q, type_=type or None)
    elif type:
        items = [item for item in alloc.catalog.all_items() if item.type == type]
    else:
        items = alloc.catalog.all_items()

    return templates.TemplateResponse(
        request,
        "_partials/inventory_rows.html",
        {"items": items},
    )


# -- pending -------------------------------------------------------------


@app.get("/pending", response_class=HTMLResponse)
def pending_page(
    request: Request,
    notified: str = "",
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
):
    if role != "admin":
        return RedirectResponse(url="/new-request", status_code=303)

    # Refresh stock readiness for all requests. This never reserves stock;
    # admin approval is still required before a request becomes Reserved.
    alloc.recheck_pending()
    recommended, analysis_rows = alloc.analyze_pending_priorities()

    return templates.TemplateResponse(
        request,
        "pending.html",
        ctx(
            request,
            role,
            active_nav="pending",
            recommended=recommended,
            analysis_rows=analysis_rows,
            notified=notified,
        ),
    )


@app.post("/pending/{project_id}/approve")
def approve_pending_request(
    project_id: str,
    role: str = Depends(require_admin),
    alloc: Allocator = Depends(get_allocator),
):
    recommended, _ = alloc.analyze_pending_priorities()
    if recommended is None:
        raise HTTPException(
            status_code=409,
            detail="No pending request currently has enough stock to reserve",
        )
    if recommended.project_id != project_id:
        raise HTTPException(
            status_code=409,
            detail=f"{recommended.project_id} has higher priority and must be handled first",
        )
    try:
        approved = alloc.approve_pending_request(project_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Project not in Pending state")
    if approved is None:
        return RedirectResponse(url="/pending", status_code=303)
    return RedirectResponse(url="/borrowed", status_code=303)


@app.post("/pending/{project_id}/notify-shortage")
def notify_pending_shortage(
    project_id: str,
    role: str = Depends(require_admin),
    alloc: Allocator = Depends(get_allocator),
):
    try:
        req = alloc.notify_stock_shortage(project_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Project not in Pending state")
    except ValueError as exc:
        raise HTTPException(status_code=409, detail=str(exc))

    return RedirectResponse(
        url=f"/pending?notified={req.project_id}",
        status_code=303,
    )


# -- reserved / borrowed / returns -----------------------------------------


@app.post("/reserved/{project_id}/confirm-handover")
def confirm_reserved_handover(
    project_id: str,
    role: str = Depends(require_admin),
    alloc: Allocator = Depends(get_allocator),
):
    try:
        alloc.confirm_handover(project_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="Project not in Reserved state")
    return RedirectResponse(url="/borrowed", status_code=303)


# -- borrowed / returns --------------------------------------------------


@app.get("/borrowed", response_class=HTMLResponse)
def borrowed_page(request: Request, role: str = Depends(get_role), alloc: Allocator = Depends(get_allocator)):
    if role != "admin":
        return RedirectResponse(url="/new-request", status_code=303)
    reserved_records = merge_sort(list(alloc.reserved.values()), key_func=lambda r: r.heap_key)
    records = merge_sort(list(alloc.borrowed.values()), key_func=lambda r: r.handed_over_at)
    return templates.TemplateResponse(
        request,
        "borrowed.html",
        ctx(
            request,
            role,
            active_nav="borrowed",
            reserved_records=reserved_records,
            records=records,
        ),
    )


@app.get("/returns", response_class=HTMLResponse)
def returns_search(
    request: Request,
    project_id: str = "",
    role: str = Depends(get_role),
    alloc: Allocator = Depends(get_allocator),
):
    record = alloc.borrowed.search(project_id) if project_id else None
    error = "Project ID not found in Borrowed" if project_id and record is None else ""
    return templates.TemplateResponse(
        request,
        "returns.html", ctx(request, role, active_nav="returns", project_id=project_id, record=record, error=error)
    )


@app.post("/returns/{project_id}")
def do_return(
    project_id: str,
    request: Request,
    role: str = Depends(require_admin),
    alloc: Allocator = Depends(get_allocator),
    item_types: list[str] = Form(default=[]),
    item_names: list[str] = Form(default=[]),
    good_quantities: list[str] = Form(default=[]),
    damaged_quantities: list[str] = Form(default=[]),
    notes: list[str] = Form(default=[]),
):
    lines: list[ReturnLineItem] = []
    for t, n, good, bad, note in zip(item_types, item_names, good_quantities, damaged_quantities, notes):
        good_q, bad_q = int(good or 0), int(bad or 0)
        if good_q > 0:
            lines.append(ReturnLineItem(t, n, good_q, ItemCondition.GOOD))
        if bad_q > 0:
            lines.append(ReturnLineItem(t, n, bad_q, ItemCondition.DAMAGED, note=note))
    try:
        alloc.process_return(project_id, lines)
    except KeyError:
        raise HTTPException(status_code=404, detail="Project not in Borrowed state")
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    return RedirectResponse(url="/borrowed", status_code=303)


# -- damaged items -------------------------------------------------------


@app.get("/damaged", response_class=HTMLResponse)
def damaged_page(request: Request, role: str = Depends(get_role), alloc: Allocator = Depends(get_allocator)):
    if role != "admin":
        return RedirectResponse(url="/new-request", status_code=303)
    records = list(alloc.damaged)
    records.reverse()
    return templates.TemplateResponse(
        request,
        "damaged.html", ctx(request, role, active_nav="damaged", records=records)
    )


# -- admin view ------------------------------------------------------------


@app.get("/admin")
def admin_view(role: str = Depends(get_role)):
    if role != "admin":
        return RedirectResponse(url="/new-request", status_code=303)
    return RedirectResponse(url="/dashboard", status_code=303)


